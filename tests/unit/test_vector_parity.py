"""``tests/eval/vector_parity.py`` -- gpu-compose 2.4's measuring tool (AC-4.9, AC-4.11).

The tool is not a test and pytest does not collect it; what is tested here is the tool: its
maths (cosine, the top-k ranking, the overlap, the Q-3 pass rule), the set it builds, and its
two network subcommands run end to end against a fake embedding service on ``127.0.0.1:0``.
No Docker, no GPU, no real model. What the tool measures on real images is design.md §5.5,
P-4 and P-5.
"""

from __future__ import annotations

import ast
import hashlib
import json
import math
import subprocess
import sys
import threading
import time
from collections.abc import Callable, Iterator, Sequence
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any

import pytest

from tests.eval import vector_parity as vp
from tests.eval.hr_handbook_set import NEGATIVES, QUESTIONS

_ROOT = Path(__file__).resolve().parents[2]
_TOOL = _ROOT / "tests" / "eval" / "vector_parity.py"


# ── helpers ─────────────────────────────────────────────────────────────────


def _unit(angle: float) -> list[float]:
    return [math.cos(angle), math.sin(angle)]


def _ids(prefix: str, count: int) -> list[str]:
    return [f"{prefix}{i}" for i in range(count)]


def _two_dimensional(
    *, question_angles: Sequence[float], chunk_angles: Sequence[float]
) -> list[list[float]]:
    return [_unit(angle) for angle in (*question_angles, *chunk_angles)]


def _evaluate(
    cpu_chunk_angles: Sequence[float],
    gpu_chunk_angles: Sequence[float],
    question_angles: Sequence[float],
    **limits: Any,
) -> dict[str, Any]:
    """Questions identical on both sides (cosine 1), chunks moved: so what a case changes is
    the RANKING, while every cosine stays far above 0.999."""
    cuts = len(question_angles)
    cpu = _two_dimensional(question_angles=question_angles, chunk_angles=cpu_chunk_angles)
    gpu = _two_dimensional(question_angles=question_angles, chunk_angles=gpu_chunk_angles)
    return vp.evaluate(
        question_ids=_ids("q", cuts),
        chunk_ids=_ids("c", len(cpu_chunk_angles)),
        cpu_questions=cpu[:cuts],
        cpu_chunks=cpu[cuts:],
        gpu_questions=gpu[:cuts],
        gpu_chunks=gpu[cuts:],
        thresholds=vp.Thresholds(**limits) if limits else None,
    )


_STEP = 1e-3
_CHUNKS = [i * _STEP for i in range(12)]  # twelve chunks, a milliradian apart


# ── cosine, top-k, overlap, fingerprint, percentile ─────────────────────────


def test_cosine_of_known_pairs() -> None:
    assert vp.cosine([1.0, 0.0], [1.0, 0.0]) == pytest.approx(1.0)
    assert vp.cosine([1.0, 0.0], [0.0, 3.0]) == pytest.approx(0.0)
    assert vp.cosine([1.0, 0.0], [-2.0, 0.0]) == pytest.approx(-1.0)
    assert vp.cosine([3.0, 4.0], [6.0, 8.0]) == pytest.approx(1.0), "scale does not matter"
    assert vp.cosine([1.0, 1.0], [1.0, 0.0]) == pytest.approx(math.sqrt(0.5))


def test_cosine_refuses_what_has_no_answer() -> None:
    with pytest.raises(ValueError, match="zero vector"):
        vp.cosine([0.0, 0.0], [1.0, 0.0])
    with pytest.raises(ValueError, match="zero vector"):
        vp.cosine([1.0, 0.0], [0.0, 0.0])
    with pytest.raises(ValueError, match="different length"):
        vp.cosine([1.0, 0.0], [1.0, 0.0, 0.0])


def test_top_k_is_by_inner_product_best_first_with_ties_to_the_lower_index() -> None:
    query = [1.0, 0.0]
    chunks = [[0.5, 0.0], [2.0, 0.0], [1.0, 0.0], [2.0, 0.0], [-1.0, 0.0]]
    assert vp.top_k(query, chunks, 3) == [1, 3, 2], "tie of chunks 1 and 3 goes to index 1"
    assert vp.top_k(query, chunks, 99) == [1, 3, 2, 0, 4], "k past the end returns them all"
    assert vp.top_k(query, chunks, 1) == [1]


def test_overlap_counts_the_shared_members() -> None:
    assert vp.overlap([1, 2, 3], [3, 2, 9]) == 2
    assert vp.overlap([1, 2], [3, 4]) == 0
    assert vp.overlap([5, 5, 5], [5]) == 1, "a set, not a multiset"


def test_fingerprint_is_ordered_framed_and_stable() -> None:
    assert vp.fingerprint(["a", "b"]) != vp.fingerprint(["b", "a"])
    assert vp.fingerprint(["ab", "c"]) != vp.fingerprint(["a", "bc"]), "framed by length"
    assert vp.fingerprint(["x"]) == vp.fingerprint(iter(["x"]))
    expected = hashlib.sha256(b"1:x\n").hexdigest()
    assert vp.fingerprint(["x"]) == expected, "the framing is part of the contract"
    assert vp.fingerprint(["س"]) == hashlib.sha256("2:س\n".encode()).hexdigest(), "UTF-8 bytes"


def test_percentile_is_nearest_rank() -> None:
    values = [5.0, 1.0, 4.0, 2.0, 3.0]
    assert vp.percentile(values, 50) == 3.0
    assert vp.percentile(values, 95) == 5.0
    assert vp.percentile(values, 100) == 5.0
    assert vp.percentile(values, 1) == 1.0
    assert vp.percentile([], 50) == 0.0


# ── the Q-3 pass rule ───────────────────────────────────────────────────────


def test_identical_vectors_pass_with_the_numbers_to_show_for_it() -> None:
    result = _evaluate(_CHUNKS, _CHUNKS, [0.0, 0.0049, 0.011])
    assert result["pass"] is True
    assert result["conditions"] == {
        "cosine_each_text": True,
        "overlap_each_question": True,
        "top1": True,
    }
    assert result["cosine"]["texts"] == 3 + 12
    assert result["cosine"]["min"] == pytest.approx(1.0)
    assert result["cosine"]["mean"] == pytest.approx(1.0)
    assert result["ranking"]["min_overlap"] == 10
    assert result["ranking"]["top1_same"] == 3
    assert result["thresholds"] == {
        "min_cosine": 0.999,
        "top_k": 10,
        "min_overlap": 9,
        "max_top1_mismatches": 1,
    }
    json.dumps(result)  # the document the tool prints must be serialisable


def test_q3_a_one_text_below_0_999_fails_the_whole_run_and_is_named() -> None:
    off = math.acos(0.9985)  # a cosine of 0.9985 between the CPU and the CUDA vector
    gpu = list(_CHUNKS)
    gpu[4] += off
    result = _evaluate(_CHUNKS, gpu, [0.0, 0.0049])
    assert result["conditions"]["cosine_each_text"] is False
    assert result["cosine"]["below_threshold"] == 1
    assert result["cosine"]["worst"][0]["text_id"] == "c4"
    assert result["cosine"]["worst"][0]["cosine"] == pytest.approx(0.9985)
    assert result["cosine"]["min"] == pytest.approx(0.9985)
    assert result["pass"] is False


def test_q3_a_is_at_least_not_more_than() -> None:
    gpu = list(_CHUNKS)
    gpu[4] += math.acos(0.9995)
    assert _evaluate(_CHUNKS, gpu, [0.0049])["conditions"]["cosine_each_text"] is True
    exact = vp.cosine(_unit(_CHUNKS[4]), _unit(gpu[4]))
    at_the_limit = _evaluate(_CHUNKS, gpu, [0.0049], min_cosine=exact)
    assert at_the_limit["conditions"]["cosine_each_text"] is True, "equal to the limit passes"
    just_above = _evaluate(_CHUNKS, gpu, [0.0049], min_cosine=exact + 1e-9)
    assert just_above["conditions"]["cosine_each_text"] is False


def test_q3_b_swapping_the_ranks_at_the_edge_of_the_top_10_still_passes_at_9() -> None:
    gpu = list(_CHUNKS)
    gpu[9], gpu[10] = gpu[10], gpu[9]  # chunk 9 drops out of the top 10 and chunk 10 comes in
    result = _evaluate(_CHUNKS, gpu, [0.0])
    assert result["ranking"]["min_overlap"] == 9
    assert result["cosine"]["min"] > 0.999, "the cosines are not what this case moves"
    assert result["conditions"]["overlap_each_question"] is True
    assert result["pass"] is True


def test_q3_b_an_overlap_of_8_fails() -> None:
    gpu = [(len(_CHUNKS) - 1 - i) * _STEP for i in range(len(_CHUNKS))]  # the ranking reversed
    result = _evaluate(_CHUNKS, gpu, [0.0])
    assert result["cosine"]["min"] > 0.999, "every cosine is fine"
    assert result["ranking"]["min_overlap"] == 8
    assert result["ranking"]["below_min_overlap"] == ["q0"]
    assert result["conditions"]["overlap_each_question"] is False
    assert result["pass"] is False


def test_q3_b_top1_tolerates_one_question_and_not_two() -> None:
    gpu = list(_CHUNKS)
    gpu[0], gpu[1] = gpu[1], gpu[0]  # the best chunk of a question at angle 0 changes
    one = _evaluate(_CHUNKS, gpu, [0.0, 0.0049, 0.0049])
    assert one["ranking"]["top1_mismatches"] == 1
    assert one["ranking"]["top1_same"] == 2
    assert one["conditions"]["top1"] is True
    assert one["pass"] is True
    two = _evaluate(_CHUNKS, gpu, [0.0, 0.0, 0.0049])
    assert two["ranking"]["top1_mismatches"] == 2
    assert two["conditions"]["top1"] is False
    assert two["conditions"]["overlap_each_question"] is True, "only the first condition failed"
    assert two["pass"] is False


def test_the_margins_are_the_ones_written_before_the_measurement() -> None:
    assert (vp.MIN_COSINE, vp.TOP_K, vp.MIN_OVERLAP, vp.MAX_TOP1_MISMATCHES) == (0.999, 10, 9, 1)
    assert vp.Thresholds() == vp.Thresholds(0.999, 10, 9, 1)


def test_evaluate_refuses_inputs_that_cannot_be_judged() -> None:
    cpu = _two_dimensional(question_angles=[0.0], chunk_angles=_CHUNKS)
    kwargs: dict[str, Any] = {
        "question_ids": ["q0"],
        "chunk_ids": _ids("c", 12),
        "cpu_questions": cpu[:1],
        "cpu_chunks": cpu[1:],
        "gpu_questions": cpu[:1],
        "gpu_chunks": cpu[1:],
    }
    assert vp.evaluate(**kwargs)["pass"] is True
    with pytest.raises(ValueError, match="too small"):
        vp.evaluate(
            **{**kwargs, "chunk_ids": _ids("c", 5), "cpu_chunks": cpu[1:6], "gpu_chunks": cpu[1:6]}
        )
    with pytest.raises(ValueError, match="question vectors"):
        vp.evaluate(**{**kwargs, "gpu_questions": []})
    with pytest.raises(ValueError, match="question vectors"):
        vp.evaluate(**{**kwargs, "question_ids": [], "cpu_questions": [], "gpu_questions": []})
    with pytest.raises(ValueError, match="chunk vectors"):
        vp.evaluate(**{**kwargs, "gpu_chunks": cpu[1:5]})
    with pytest.raises(ValueError, match="zero vector"):
        vp.evaluate(**{**kwargs, "gpu_questions": [[0.0, 0.0]]})


# ── the set ─────────────────────────────────────────────────────────────────


def test_the_42_question_texts_are_the_21_questions_in_two_languages() -> None:
    items = vp.question_items(QUESTIONS, NEGATIVES)
    assert len(QUESTIONS) == 15 and len(NEGATIVES) == 6
    assert len(items) == 42
    assert len({item["id"] for item in items}) == 42
    assert sorted(item["lang"] for item in items) == ["ar"] * 21 + ["en"] * 21
    assert items[0] == {"id": "q1-en", "lang": "en", "text": QUESTIONS[0]["en"]}
    assert items[1] == {"id": "q1-ar", "lang": "ar", "text": QUESTIONS[0]["ar"]}
    assert items[-1]["id"] == "n6-ar"
    assert all(item["text"].strip() for item in items)


def test_soak_texts_are_arabic_350_words_deterministic_and_different() -> None:
    corpus = vp._arabic_corpus()
    assert len(corpus) == 21, "the 15 questions and the 6 unanswerable ones, in Arabic"
    words = " ".join(corpus).split()
    first = vp.soak_text(words, 0)
    assert len(first.split()) == 350
    assert first == vp.soak_text(words, 0), "a function of its arguments alone"
    assert vp.soak_text(words, 1) != first
    assert len({vp.soak_text(words, i) for i in range(64)}) == 64
    assert set(first.split()) <= set(words)
    arabic = sum("؀" <= ch <= "ۿ" for ch in first)
    assert arabic > len(first) * 0.7, "Arabic text, as AC-4.11 measures"
    assert len(vp.soak_text(words, 3, length=10).split()) == 10
    with pytest.raises(ValueError, match="no words"):
        vp.soak_text([], 0)


def _write_set(path: Path, questions: int = 5, chunks: int = 14) -> dict[str, Any]:
    items = [
        {"id": f"q{i}", "lang": "en", "text": f"question number {i}"} for i in range(questions)
    ]
    pieces = [{"id": f"chunk-{i:04d}", "text": f"chunk of the handbook {i}"} for i in range(chunks)]
    data = {
        "version": vp.SET_VERSION,
        "fingerprint": vp.fingerprint([x["text"] for x in (*items, *pieces)]),
        "document": {"source": "test", "sha256": "0" * 64},
        "chunking": {"max_words": 192},
        "questions": items,
        "chunks": pieces,
    }
    path.write_text(json.dumps(data), encoding="utf-8")
    return data


def test_load_set_accepts_what_it_wrote_and_refuses_an_edited_text(tmp_path: Path) -> None:
    path = tmp_path / "set.json"
    data = _write_set(path)
    assert vp.load_set(path)["fingerprint"] == data["fingerprint"]
    data["chunks"][3]["text"] = "edited by hand"
    path.write_text(json.dumps(data), encoding="utf-8")
    with pytest.raises(vp.ToolError, match="fingerprint does not match"):
        vp.load_set(path)


@pytest.mark.parametrize(
    "content",
    ["not json", "[]", '{"version": 99}', '{"version": 1}', '{"version": 1, "questions": 1}'],
)
def test_load_set_refuses_what_is_not_a_set(tmp_path: Path, content: str) -> None:
    path = tmp_path / "set.json"
    path.write_text(content, encoding="utf-8")
    with pytest.raises(vp.ToolError):
        vp.load_set(path)
    with pytest.raises(vp.ToolError):
        vp.load_set(tmp_path / "missing.json")


def _git(repo: Path, *args: str) -> str:
    done = subprocess.run(
        ["git", "-c", "user.name=qa", "-c", "user.email=qa@example.invalid", *args],
        cwd=repo,
        capture_output=True,
        text=True,
        check=True,
    )
    return done.stdout.strip()


def test_the_handbook_comes_from_the_tree_else_from_the_commit_that_added_it(
    tmp_path: Path,
) -> None:
    repo = tmp_path / "repo"
    (repo / "docs").mkdir(parents=True)
    _git(repo, "init", "-q")
    handbook = repo / vp.DEFAULT_DOCUMENT
    handbook.write_bytes(b"the handbook v1")
    _git(repo, "add", ".")
    _git(repo, "commit", "-q", "--no-gpg-sign", "-m", "add")
    added = _git(repo, "rev-parse", "HEAD")
    handbook.unlink()
    _git(repo, "commit", "-q", "--no-gpg-sign", "-am", "remove")

    data, source = vp.read_document(None, repo=repo, rev=added)
    assert data == b"the handbook v1"
    assert source == {
        "source": f"git:{added}:{vp.DEFAULT_DOCUMENT}",
        "sha256": hashlib.sha256(b"the handbook v1").hexdigest(),
    }
    handbook.write_bytes(b"the handbook v2")
    data, source = vp.read_document(None, repo=repo, rev=added)
    assert (data, source["source"]) == (b"the handbook v2", vp.DEFAULT_DOCUMENT), "the tree wins"
    other = tmp_path / "other.docx"
    other.write_bytes(b"other")
    data, source = vp.read_document(other, repo=repo, rev=added)
    assert data == b"other" and source["source"] == str(other)
    with pytest.raises(vp.ToolError, match="does not exist"):
        vp.read_document(tmp_path / "nope.docx", repo=repo, rev=added), "no silent fallback"
    handbook.unlink()
    with pytest.raises(vp.ToolError, match="git show"):
        vp.read_document(None, repo=repo, rev="0" * 40)


def _small_docx(path: Path) -> None:
    from docx import Document  # noqa: PLC0415

    document = Document()
    document.add_heading("Working hours", level=1)
    document.add_paragraph(
        "Employees work forty-eight hours per week. A delay of up to ten minutes is allowed."
    )
    document.add_heading("Leave", level=1)
    document.add_paragraph("Sick leave needs a medical report within twenty-four hours.")
    table = document.add_table(rows=3, cols=2)
    for row, (item, detail) in enumerate(
        [
            ("Item", "Details"),
            ("Insurance", "Spouse and children"),
            ("Overtime", "One and a quarter on weekdays"),
        ]
    ):
        table.cell(row, 0).text = item
        table.cell(row, 1).text = detail
    document.save(str(path))


def test_build_set_uses_the_platforms_parser_chunker_and_table_explosion(tmp_path: Path) -> None:
    docx = tmp_path / "handbook.docx"
    _small_docx(docx)
    built = vp.build_set(document=docx)
    assert [q["id"] for q in built["questions"]] == [
        item["id"] for item in vp.question_items(QUESTIONS, NEGATIVES)
    ]
    assert len(built["questions"]) == 42
    texts = [c["text"] for c in built["chunks"]]
    assert texts, "the document produced chunks"
    assert [c["id"] for c in built["chunks"]] == [f"chunk-{i:04d}" for i in range(len(texts))]
    joined = "\n".join(texts)
    assert "forty-eight hours per week" in joined
    assert "Spouse and children" in joined, "a table row reached a chunk"
    assert built["chunking"]["max_words"] == 192, "the platform's window at 512 tokens"
    assert built["chunking"]["table_blocks"] == 1
    assert built["chunking"]["semantic_pre_splitting"] is False
    assert built["document"]["sha256"] == hashlib.sha256(docx.read_bytes()).hexdigest()
    assert built["fingerprint"] == vp.fingerprint(
        [*(q["text"] for q in built["questions"]), *texts]
    )
    again = vp.build_set(document=docx)
    assert again == built, "the same document gives the same set"


def test_build_set_writes_a_file_compare_can_read(tmp_path: Path) -> None:
    docx = tmp_path / "handbook.docx"
    _small_docx(docx)
    out = tmp_path / "set.json"
    assert vp.main(["build-set", "--out", str(out), "--document", str(docx)]) == 0
    assert vp.load_set(out)["document"]["source"] == str(docx)
    assert vp.main(["build-set", "--out", str(out), "--document", str(tmp_path / "no.docx")]) == 3


# ── a fake embedding service ────────────────────────────────────────────────


def _fake_vector(text: str, *, noise: float, dim: int = 16) -> list[float]:
    base = [b - 127.5 for b in hashlib.sha256(text.encode("utf-8")).digest()[:dim]]
    if noise:
        jitter = hashlib.sha256(b"noise:" + text.encode("utf-8")).digest()[:dim]
        base = [x + noise * (j - 127.5) for x, j in zip(base, jitter, strict=True)]
    norm = math.sqrt(sum(x * x for x in base))
    return [x / norm for x in base]


class FakeService:
    """``GET /health`` and ``POST /embed`` of the embedding service, with the vectors of a
    hash. ``noise`` is how far this "image" strays from the exact vector."""

    def __init__(
        self,
        *,
        device: str | None = "cpu",
        noise: float = 0.0,
        model: str = "fake/model",
        embed_status: int = 200,
        delay_s: float = 0.0,
    ) -> None:
        self.embeds: list[list[str]] = []
        self.active = 0
        self.max_active = 0
        self._lock = threading.Lock()
        outer = self
        health: dict[str, Any] = {
            "status": "ok",
            "model": model,
            "dimensions": 16,
            "max_seq_length": 512,
            "dtype": "float16" if device == "cuda" else "float32",
        }
        if device is not None:
            health["device"] = device

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args: object) -> None:
                return

            def _send(self, status: int, payload: dict[str, Any]) -> None:
                body = json.dumps(payload).encode()
                self.send_response(status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def do_GET(self) -> None:
                if self.path == "/health":
                    self._send(200, health)
                else:
                    self._send(404, {"detail": "no"})

            def do_POST(self) -> None:
                body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
                if self.path != "/embed":
                    self._send(404, {"detail": "no"})
                    return
                with outer._lock:
                    outer.embeds.append(list(body["texts"]))
                    outer.active += 1
                    outer.max_active = max(outer.max_active, outer.active)
                try:
                    if delay_s:
                        time.sleep(delay_s)
                    if embed_status != 200:
                        self._send(embed_status, {"detail": "scripted"})
                    elif body["model"] != model:
                        self._send(400, {"detail": "model mismatch"})
                    else:
                        vectors = [_fake_vector(t, noise=noise) for t in body["texts"]]
                        self._send(
                            200,
                            {"vectors": vectors, "model": model, "dimensions": 16, "tokens": 1},
                        )
                finally:
                    with outer._lock:
                        outer.active -= 1

        self.server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        self._thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self._thread.start()

    @property
    def url(self) -> str:
        host, port = self.server.server_address[:2]
        return f"http://{host!s}:{port}"

    def close(self) -> None:
        self.server.shutdown()
        self.server.server_close()


@pytest.fixture
def service() -> Iterator[Callable[..., FakeService]]:
    made: list[FakeService] = []

    def make(**kwargs: Any) -> FakeService:
        instance = FakeService(**kwargs)
        made.append(instance)
        return instance

    yield make
    for instance in made:
        instance.close()


def _compare(
    capsys: pytest.CaptureFixture[str],
    cpu: FakeService,
    gpu: FakeService,
    set_path: Path,
    *extra: str,
) -> tuple[int, dict[str, Any] | None, str]:
    code = vp.main(["compare", "--cpu", cpu.url, "--gpu", gpu.url, "--set", str(set_path), *extra])
    captured = capsys.readouterr()
    return code, (json.loads(captured.out) if captured.out.strip() else None), captured.err


# ── compare, end to end ─────────────────────────────────────────────────────


def test_compare_passes_for_a_cuda_image_that_agrees_and_prints_the_evidence(
    tmp_path: Path, service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    set_path = tmp_path / "set.json"
    data = _write_set(set_path)
    cpu = service(device="cpu")
    gpu = service(device="cuda", noise=1e-3)
    code, result, err = _compare(capsys, cpu, gpu, set_path)
    assert code == 0, err
    assert result is not None and result["pass"] is True
    assert 0.999 <= result["cosine"]["min"] < 1.0, "a real, nonzero difference was measured"
    assert result["ranking"]["questions"] == 5
    assert result["set"]["fingerprint"] == data["fingerprint"]
    assert (result["set"]["questions"], result["set"]["chunks"]) == (5, 14)
    assert result["services"]["cpu"]["health"]["device"] == "cpu"
    assert result["services"]["gpu"]["health"]["dtype"] == "float16"
    sent = {text for batch in cpu.embeds for text in batch}
    assert sent == {x["text"] for x in (*data["questions"], *data["chunks"])}
    assert cpu.embeds == gpu.embeds, "both saw the same batches"


def test_compare_exits_nonzero_when_the_cuda_vectors_stray(
    tmp_path: Path, service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    set_path = tmp_path / "set.json"
    _write_set(set_path)
    code, result, _ = _compare(
        capsys, service(device="cpu"), service(device="cuda", noise=0.3), set_path
    )
    assert code == 1
    assert result is not None and result["pass"] is False
    assert result["conditions"]["cosine_each_text"] is False
    assert result["cosine"]["below_threshold"] > 0
    assert result["cosine"]["worst"], "the worst texts are named for the report"


def test_compare_will_not_call_two_images_on_one_device_a_comparison(
    tmp_path: Path, service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    set_path = tmp_path / "set.json"
    _write_set(set_path)
    cpu, same = service(device="cpu"), service(device="cpu")
    code, result, err = _compare(capsys, cpu, same, set_path)
    assert (code, result) == (3, None)
    assert "not a CPU-against-CUDA comparison" in err
    assert cpu.embeds == [], "nothing was embedded"
    swapped = _compare(capsys, service(device="cuda"), service(device="cuda"), set_path)
    assert swapped[0] == 3
    code, result, _ = _compare(capsys, cpu, same, set_path, "--allow-same-device")
    assert code == 0 and result is not None and result["cosine"]["min"] == pytest.approx(1.0)


def test_compare_refuses_an_image_that_does_not_say_what_it_ran_on(
    tmp_path: Path, service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    set_path = tmp_path / "set.json"
    _write_set(set_path)
    code, _, err = _compare(capsys, service(device="cpu"), service(device=None), set_path)
    assert code == 3
    assert "does not report `device`" in err


def test_compare_refuses_two_services_that_are_not_the_same_model(
    tmp_path: Path, service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    set_path = tmp_path / "set.json"
    _write_set(set_path)
    code, _, err = _compare(
        capsys, service(device="cpu"), service(device="cuda", model="other/model"), set_path
    )
    assert code == 3
    assert "differ in model" in err


def test_compare_refuses_an_unreachable_service_and_an_edited_set(
    tmp_path: Path, service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    set_path = tmp_path / "set.json"
    data = _write_set(set_path)
    live = service(device="cpu")
    dead = service(device="cuda")
    dead.close()
    code, _, err = _compare(capsys, live, dead, set_path)
    assert code == 3 and "failed" in err
    broken = service(device="cuda", embed_status=500)
    code, _, err = _compare(capsys, live, broken, set_path)
    assert code == 3 and "HTTP 500" in err
    data["questions"][0]["text"] = "edited"
    set_path.write_text(json.dumps(data), encoding="utf-8")
    code, _, err = _compare(capsys, live, service(device="cuda"), set_path)
    assert code == 3 and "fingerprint" in err


def test_compare_refuses_a_set_too_small_to_rank(
    tmp_path: Path, service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    set_path = tmp_path / "set.json"
    _write_set(set_path, questions=3, chunks=6)
    code, _, err = _compare(capsys, service(device="cpu"), service(device="cuda"), set_path)
    assert code == 3 and "too small" in err


# ── soak, end to end ────────────────────────────────────────────────────────


def test_soak_sends_the_ac_4_11_load_and_reports_it(
    service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    svc = service(device="cuda", delay_s=0.05)
    code = vp.main(
        ["soak", "--url", svc.url, "--minutes", "0.03", "--concurrency", "3", "--batch", "2"]
    )
    captured = capsys.readouterr()
    result = json.loads(captured.out)
    assert code == 0, captured.err
    assert result["requests_ok"] == len(svc.embeds) > 3
    assert result["requests_failed"] == 0 and result["errors"] == {}
    assert result["texts_embedded"] == result["requests_ok"] * 2
    assert all(len(batch) == 2 for batch in svc.embeds), "the batch size asked for"
    assert all(len(text.split()) == 350 for batch in svc.embeds for text in batch)
    assert svc.max_active >= 2, "concurrent calls, not a loop of single ones"
    assert result["latency_s"]["p50"] >= 0.05
    assert result["health_before"]["device"] == "cuda"
    assert result["interrupted"] is False
    assert result["duration_s"] >= 0.03 * 60


def test_soak_texts_differ_between_requests(
    service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    svc = service(device="cpu")
    code = vp.main(
        ["soak", "--url", svc.url, "--minutes", "0.01", "--concurrency", "1", "--batch", "3"]
    )
    capsys.readouterr()
    assert code == 0
    texts = [text for batch in svc.embeds for text in batch]
    assert len(texts) > 3 and len(set(texts)) > 3


def test_soak_exits_nonzero_when_the_service_errors(
    service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    svc = service(device="cpu", embed_status=500)
    code = vp.main(["soak", "--url", svc.url, "--minutes", "0.01", "--concurrency", "2"])
    result = json.loads(capsys.readouterr().out)
    assert code == 1
    assert result["requests_ok"] == 0 and result["requests_failed"] > 0
    assert any("HTTP 500" in message for message in result["errors"])


def test_soak_against_a_service_that_is_down_is_unusable(
    service: Callable[..., FakeService], capsys: pytest.CaptureFixture[str]
) -> None:
    svc = service(device="cpu")
    svc.close()
    assert vp.main(["soak", "--url", svc.url, "--minutes", "0.01"]) == 3
    assert "ERROR" in capsys.readouterr().err


# ── the command line, and what kind of file this is ─────────────────────────


@pytest.mark.parametrize(
    "argv",
    [
        [],
        ["frobnicate"],
        ["soak"],
        ["soak", "--url", "http://x", "--minutes", "0"],
        ["soak", "--url", "http://x", "--minutes", "inf"],
        ["soak", "--url", "http://x", "--concurrency", "0"],
        ["compare", "--cpu", "http://x", "--gpu", "http://y"],
        ["build-set"],
    ],
)
def test_bad_usage_is_exit_2(argv: list[str], capsys: pytest.CaptureFixture[str]) -> None:
    assert vp.main(argv) == 2
    assert capsys.readouterr().err


def test_the_defaults_are_the_ac_4_11_load() -> None:
    args = vp._parser().parse_args(["soak", "--url", "http://x"])
    assert (args.minutes, args.concurrency, args.batch) == (15.0, 8, 8)
    assert vp.SOAK_WORDS == 350


def _top_level_imports(tree: ast.Module) -> set[str]:
    roots: set[str] = set()
    for node in tree.body:
        if isinstance(node, ast.Import):
            roots |= {alias.name.split(".")[0] for alias in node.names}
        elif isinstance(node, ast.ImportFrom) and node.module and node.level == 0:
            roots.add(node.module.split(".")[0])
    return roots


def test_compare_and_soak_need_the_standard_library_only() -> None:
    """They run in a stock python:3.12-slim container (P-5): only `build-set`, which needs the
    platform's parser, may import anything else, and only inside its own function."""
    tree = ast.parse(_TOOL.read_text(encoding="utf-8"))
    roots = _top_level_imports(tree)
    assert roots, "the import scan matched nothing"
    assert roots <= set(sys.stdlib_module_names), sorted(roots - set(sys.stdlib_module_names))
    inner = {
        alias.name.split(".")[0]
        for fn in ast.walk(tree)
        if isinstance(fn, ast.FunctionDef)
        for node in ast.walk(fn)
        if isinstance(node, ast.Import)
        for alias in node.names
    } | {
        node.module.split(".")[0]
        for fn in ast.walk(tree)
        if isinstance(fn, ast.FunctionDef)
        for node in ast.walk(fn)
        if isinstance(node, ast.ImportFrom) and node.module and node.level == 0
    }
    assert inner - set(sys.stdlib_module_names) == {"app"}, "only the platform, only in a function"


def test_the_tool_is_not_a_test_and_pytest_collects_nothing_from_it() -> None:
    assert not _TOOL.name.startswith("test_") and not _TOOL.name.endswith("_test.py")
    tree = ast.parse(_TOOL.read_text(encoding="utf-8"))
    names = [n.name for n in ast.walk(tree) if isinstance(n, ast.ClassDef | ast.FunctionDef)]
    assert not [n for n in names if n.startswith(("test_", "Test"))], names
    done = subprocess.run(
        [
            sys.executable,
            "-m",
            "pytest",
            "--collect-only",
            "-q",
            "-p",
            "no:cacheprovider",
            str(_TOOL.parent),
        ],
        cwd=_ROOT,
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    assert done.returncode == 5, done.stdout + done.stderr  # "no tests collected"
    assert "vector_parity" not in done.stdout


def test_run_as_a_script_it_prints_its_usage() -> None:
    done = subprocess.run(
        [sys.executable, "-I", str(_TOOL), "soak", "--help"],
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert done.returncode == 0, done.stderr
    assert "--minutes" in done.stdout and "--concurrency" in done.stdout
